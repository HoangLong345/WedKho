from __future__ import annotations

import re
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from app.domain.pallet_calculator import ItemCalculationError, ItemInput


class ExcelImportError(ValueError):
    pass


_REQUIRED_HEADERS = {
    "item_code": (
        "mã kiện",
        "ma kien",
        "mã kiện hàng",
        "ma kien hang",
        "t xếp pallet - mã kiện",
        "mã kiện/pallet",
        "kiện hàng",
        "mã pallet",
        "ma pallet",
        "mã hàng",
        "mã hàng hóa",
        "ma hang",
        "ma hang hoa",
        "mã sku",
        "ma sku",
        "sku",
        "item_code",
        "item code",
        "mã sp",
        "ma sp",
        "mã sản phẩm",
        "code",
        "mã",
    ),
    "item_name": (
        "tên hàng",
        "tên hàng hóa",
        "ten hang",
        "ten hang hoa",
        "tên sku",
        "tên sản phẩm",
        "ten san pham",
        "tên sp",
        "item_name",
        "item name",
        "name",
        "mặt hàng",
        "mat hang",
        "tên mặt hàng",
        "diễn giải",
        "mô tả",
    ),
    "quantity": (
        "số lượng",
        "so luong",
        "quantity",
        "qty",
        "số lượng kiện",
        "so luong kien",
        "số kiện",
        "so kien",
        "kiện",
        "kien",
        "sl",
        "sl kiện",
        "số lượng (kiện)",
        "tổng số kiện",
    ),
    "dimensions": (
        "quy cách 1 kiện",
        "quy cach 1 kien",
        "quy cách",
        "quy cach",
        "kích thước kiện d × r × c (mm)",
        "kích thước kiện d x r x c (mm)",
        "kích thước kiện d × r × c",
        "kích thước kiện d x r x c",
        "kích thước kiện",
        "kich thuoc kien",
        "kích thước",
        "kich thuoc",
        "d × r × c",
        "d x r x c",
        "d*r*c",
        "dxrxc",
        "dimensions",
        "dimension",
        "quy cách đóng gói",
        "quy cach dong goi",
        "kích thước (mm)",
        "kích thước (cm)",
        "kích thước (m)",
        "kích thước kiện (mm)",
    ),
    "gross_weight_kg": (
        "g.w. mỗi kiện (kg)",
        "g.w. mỗi kiện",
        "gw mỗi kiện",
        "gw moi kien",
        "g.w mỗi kiện",
        "g.w/kiện",
        "g.w/kiện (kg)",
        "gw/kiện",
        "gw/kien",
        "gw/kiện (kg)",
        "gross weight",
        "gross_weight_kg",
        "gross weight (kg)",
        "khối lượng mỗi kiện (kg)",
        "khối lượng mỗi kiện",
        "khối lượng/kiện",
        "trọng lượng mỗi kiện (kg)",
        "trọng lượng mỗi kiện",
        "trọng lượng/kiện",
        "trọng lượng kiện",
        "khối lượng kiện",
        "g.w (kg)",
        "gw (kg)",
        "g.w",
        "gw",
    ),
}


def _normalize_header(value: Any) -> str:
    return " ".join(str(value or "").strip().lower().replace("_", " ").split())


def _find_header(headers: dict[str, int], aliases: tuple[str, ...]) -> int | None:
    for alias in aliases:
        if _normalize_header(alias) in headers:
            return headers[_normalize_header(alias)]
    return None


def _find_header_containing(headers: dict[str, int], terms: tuple[str, ...]) -> int | None:
    for header, index in headers.items():
        if all(term in header for term in terms):
            return index
    return None


def _number(value: Any, row_number: int, label: str) -> float:
    try:
        cleaned = str(value).strip().replace(",", ".")
        number = float(cleaned)
    except (TypeError, ValueError) as exc:
        raise ExcelImportError(
            f"Dòng {row_number}: {label} phải là số"
        ) from exc
    if number <= 0:
        raise ExcelImportError(f"Dòng {row_number}: {label} phải lớn hơn 0")
    return number


def _clean_dimension_part(part_str: str, is_mm: bool, row_number: int) -> float:
    s = part_str.strip()
    if is_mm:
        # Nếu có dấu chấm hàng nghìn như 2.438 hoặc 1.016 hoặc 2.438,5
        if re.match(r"^\d{1,3}(?:\.\d{3})+(?:,\d+)?$", s):
            s = s.replace(".", "").replace(",", ".")
        elif re.match(r"^\d{1,3}\.\d{3}$", s):
            s = s.replace(".", "")
        else:
            s = s.replace(",", ".")
    else:
        s = s.replace(",", ".")
    return _number(s, row_number, "kích thước")


def _dimensions(
    value: Any, row_number: int, unit_hint: str | None = None
) -> tuple[float, float, float, str | None]:
    raw = (
        str(value or "")
        .lower()
        .replace("×", "x")
        .replace("*", "x")
    )
    detected_unit = None
    if "mm" in raw:
        detected_unit = "mm"
        raw = raw.replace("mm", "")
    elif "cm" in raw:
        detected_unit = "cm"
        raw = raw.replace("cm", "")
    elif "m" in raw:
        detected_unit = "m"
        raw = raw.replace("m", "")

    parts = [part.strip() for part in raw.split("x")]
    if len(parts) != 3:
        raise ExcelImportError(
            f"Dòng {row_number}: quy cách phải có dạng D×R×C, ví dụ 60x40x30"
        )

    effective_unit = detected_unit or unit_hint
    is_mm = effective_unit == "mm"

    # Nếu chưa xác định rõ đơn vị nhưng thấy số dạng phân cách nghìn chấm (VD: 2.438) hoặc số >= 100
    if not is_mm and effective_unit is None:
        has_thousand_dot = any(re.match(r"^\d{1,3}\.\d{3}$", p) for p in parts)
        has_large_number = any(
            p.replace(",", ".").replace(".", "").isdigit()
            and int(p.replace(",", ".").replace(".", "")) >= 100
            for p in parts
        )
        if has_thousand_dot or has_large_number:
            is_mm = True
            detected_unit = "mm"

    dims = tuple(_clean_dimension_part(part, is_mm, row_number) for part in parts)
    return dims[0], dims[1], dims[2], detected_unit


def parse_excel_inbound(workbook_bytes: bytes) -> dict[str, Any]:
    workbook = load_workbook(BytesIO(workbook_bytes), read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    try:
        raw_headers = next(rows)
    except StopIteration as exc:
        raise ExcelImportError("File Excel không có dòng tiêu đề") from exc

    headers = {
        _normalize_header(value): index
        for index, value in enumerate(raw_headers)
        if value is not None
    }
    columns = {
        name: _find_header(headers, aliases)
        for name, aliases in _REQUIRED_HEADERS.items()
    }

    # Fallbacks thông minh cho item_code và item_name
    if columns["item_code"] is None:
        for h, idx in headers.items():
            if any(k in h for k in ("mã kiện", "ma kien", "mã hàng", "ma hang", "sku", "item code", "item_code", "kiện")):
                columns["item_code"] = idx
                break

    if columns["item_name"] is None:
        for h, idx in headers.items():
            if any(k in h for k in ("tên hàng", "ten hang", "tên sp", "ten sp", "tên sản phẩm", "mặt hàng")):
                columns["item_name"] = idx
                break

    # Nếu file chỉ có cột mã hoặc cột tên -> dùng chung cho cả hai
    if columns["item_name"] is None and columns["item_code"] is not None:
        columns["item_name"] = columns["item_code"]
    elif columns["item_code"] is None and columns["item_name"] is not None:
        columns["item_code"] = columns["item_name"]

    # Fallback cho quantity (số lượng / số kiện)
    if columns["quantity"] is None:
        for h, idx in headers.items():
            if any(k in h for k in ("số kiện", "so kien", "số lượng", "so luong", "quantity", "qty")):
                if (
                    "kích thước" not in h
                    and "quy cách" not in h
                    and "g.w" not in h
                    and "khối lượng" not in h
                    and "trọng lượng" not in h
                ):
                    columns["quantity"] = idx
                    break

    separate_dimension_columns = {
        "length": _find_header(headers, ("dài", "dai", "chiều dài", "length", "length_m")),
        "width": _find_header(headers, ("rộng", "rong", "chiều rộng", "width", "width_m")),
        "height": _find_header(headers, ("cao", "chiều cao", "height", "height_m")),
    }
    has_separate_dimensions = all(
        index is not None for index in separate_dimension_columns.values()
    )

    # Fallback cho dimensions (nếu chưa có 3 cột Dài/Rộng/Cao riêng)
    if has_separate_dimensions:
        columns["dimensions"] = None
    elif columns["dimensions"] is None:
        for h, idx in headers.items():
            if "đơn vị" in h or "don vi" in h or "unit" in h:
                continue
            if any(
                k in h
                for k in (
                    "kích thước",
                    "kich thuoc",
                    "quy cách",
                    "quy cach",
                    "d × r × c",
                    "d x r x c",
                    "d*r*c",
                    "dimensions",
                )
            ):
                columns["dimensions"] = idx
                break

    # Fallback cho gross_weight_kg (tránh nhầm với tổng khối lượng hoặc khối lượng mã)
    if columns["gross_weight_kg"] is None:
        for h, idx in headers.items():
            if "mã" in h or "tổng" in h:
                continue
            if "g.w" in h or "gw" in h:
                columns["gross_weight_kg"] = idx
                break
            if ("khối lượng" in h or "trọng lượng" in h) and ("kiện" in h or "mỗi kiện" in h):
                columns["gross_weight_kg"] = idx
                break
        if columns["gross_weight_kg"] is None:
            for h, idx in headers.items():
                if "mã" in h or "tổng" in h:
                    continue
                if ("khối lượng" in h or "trọng lượng" in h) and "kg" in h:
                    columns["gross_weight_kg"] = idx
                    break

    dimensions_available = columns["dimensions"] is not None or has_separate_dimensions
    column_label_vi = {
        "item_code": "Mã SKU / Mã hàng",
        "item_name": "Tên mặt hàng",
        "quantity": "Số lượng kiện",
        "gross_weight_kg": "Khối lượng kiện (kg)",
    }
    missing = [
        column_label_vi.get(name, name)
        for name, index in columns.items()
        if index is None and name not in ("dimensions", "item_name")
    ]
    if columns["item_name"] is None and columns["item_code"] is None:
        missing.append("Mã hàng hoặc Tên hàng")
    if not dimensions_available:
        missing.append("Kích thước kiện (Dài x Rộng x Cao hoặc 3 cột riêng Dài, Rộng, Cao)")
    if missing:
        raise ExcelImportError(
            "Tệp Excel thiếu các cột bắt buộc: " + ", ".join(missing)
        )

    optional = {
        "item_type": ("loại hàng", "loai hang", "item_type", "category"),
        "unit_of_measure": ("đvt", "unit", "unit_of_measure"),
        "dimension_unit": ("đơn vị kích thước", "don vi kich thuoc", "dimension_unit"),
        "tenant_id": (
            "khách hàng (viết tắt)",
            "khach hang (viet tat)",
            "mã khách hàng",
            "ma khach hang",
            "mã đối tác",
            "tenant_id",
            "tenant code",
        ),
        "tenant_name": (
            "khách hàng",
            "khach hang",
            "tên khách hàng",
            "ten khach hang",
            "tên đối tác",
            "tenant_name",
            "tenant",
        ),
    }
    optional_columns = {
        name: _find_header(headers, aliases) for name, aliases in optional.items()
    }

    # Chỉ tìm đơn vị trên các cột liên quan đến kích thước / quy cách
    dimension_headers = []
    for h, idx in headers.items():
        if idx == columns.get("dimensions"):
            dimension_headers.append(h)
        elif idx in separate_dimension_columns.values():
            dimension_headers.append(h)
        elif (
            optional_columns.get("dimension_unit") is not None
            and idx == optional_columns["dimension_unit"]
        ):
            dimension_headers.append(h)

    header_unit: str | None = None
    for h in dimension_headers:
        if "(mm)" in h or "[mm]" in h or " mm" in h or h.endswith("mm"):
            header_unit = "mm"
            break
        elif "(cm)" in h or "[cm]" in h or " cm" in h or h.endswith("cm"):
            header_unit = "cm"
            break
        elif "(m)" in h or "[m]" in h or " m" in h or h.endswith(" m") or h.endswith("(m)"):
            header_unit = "m"
            break

    items: list[ItemInput] = []
    detected_tenant_code: str | None = None
    detected_tenant_name: str | None = None
    oversized_items: list[dict[str, Any]] = []

    for row_number, row in enumerate(rows, start=2):
        if not any(value is not None and str(value).strip() for value in row):
            continue
        item_code = str(row[columns["item_code"]] or "").strip()
        item_name = str(row[columns["item_name"]] or "").strip()

        norm_code = _normalize_header(item_code)
        norm_name = _normalize_header(item_name)
        if any(norm_code.startswith(prefix) for prefix in ("tổng", "tong", "total", "sum")) or \
           any(norm_name.startswith(prefix) for prefix in ("tổng", "tong", "total", "sum")):
            continue

        if not item_code and not item_name:
            continue

        # Đọc tenant code và tenant name nếu chưa có
        if not detected_tenant_code and optional_columns.get("tenant_id") is not None:
            val = str(row[optional_columns["tenant_id"]] or "").strip()
            if val and not any(val.lower().startswith(p) for p in ("tổng", "total")):
                detected_tenant_code = val

        if not detected_tenant_name and optional_columns.get("tenant_name") is not None:
            val = str(row[optional_columns["tenant_name"]] or "").strip()
            if val and not any(val.lower().startswith(p) for p in ("tổng", "total")):
                detected_tenant_name = val

        # Nếu dùng chung 1 cột và có định dạng 'Mã: Tên' thì tự tách cho trực quan
        if columns["item_code"] == columns["item_name"]:
            if ":" in item_code:
                code_part, name_part = [p.strip() for p in item_code.split(":", 1)]
                if code_part and name_part:
                    item_code = code_part
                    item_name = name_part
            elif not item_name:
                item_name = item_code

        if not item_code:
            item_code = item_name

        if columns["dimensions"] is not None:
            length, width, height, detected_unit = _dimensions(
                row[columns["dimensions"]], row_number, unit_hint=header_unit
            )
        else:
            length = _number(row[separate_dimension_columns["length"]], row_number, "Dài")
            width = _number(row[separate_dimension_columns["width"]], row_number, "Rộng")
            height = _number(row[separate_dimension_columns["height"]], row_number, "Cao")
            detected_unit = None

        explicit_unit = (
            str(row[optional_columns["dimension_unit"]] or "").strip().lower()
            if optional_columns["dimension_unit"] is not None and row[optional_columns["dimension_unit"]]
            else None
        )
        dimension_unit = (
            explicit_unit
            or detected_unit
            or header_unit
            or _infer_dimension_unit((length, width, height))
        )
        if dimension_unit not in {"mm", "cm", "m"}:
            raise ExcelImportError(
                f"Dòng {row_number}: đơn vị kích thước phải là mm, cm hoặc m"
            )

        u_factor = {"m": 1.0, "cm": 0.01, "mm": 0.001}.get(dimension_unit, 1.0)
        len_m = length * u_factor
        wid_m = width * u_factor
        hgt_m = height * u_factor
        gw_kg = _number(row[columns["gross_weight_kg"]], row_number, "G.W/kiện")

        # Kiểm tra quá khổ khoang kệ dầm beam 3.8m x 1.8m hoặc quá 4500kg
        max_dim = max(len_m, wid_m)
        min_dim = min(len_m, wid_m)
        if max_dim > 3.8:
            oversized_items.append({
                "item_code": item_code,
                "item_name": item_name,
                "length_m": len_m,
                "width_m": wid_m,
                "height_m": hgt_m,
                "gross_weight_kg": gw_kg,
                "reason": f"Chiều dài {max_dim:.2f}m vượt quá kích thước lọt lòng khoang kệ (tối đa 3.8m)",
            })
        elif min_dim > 1.8:
            oversized_items.append({
                "item_code": item_code,
                "item_name": item_name,
                "length_m": len_m,
                "width_m": wid_m,
                "height_m": hgt_m,
                "gross_weight_kg": gw_kg,
                "reason": f"Chiều rộng {min_dim:.2f}m vượt quá kích thước lọt lòng khoang kệ (tối đa 1.8m)",
            })
        elif gw_kg > 4500:
            oversized_items.append({
                "item_code": item_code,
                "item_name": item_name,
                "length_m": len_m,
                "width_m": wid_m,
                "height_m": hgt_m,
                "gross_weight_kg": gw_kg,
                "reason": f"Khối lượng kiện ({gw_kg:.0f}kg) vượt quá tải trọng tối đa Tầng 1 (4.500kg)",
            })

        items.append(
            ItemInput(
                item_code=item_code,
                item_name=item_name,
                item_type=(
                    str(row[optional_columns["item_type"]] or "").strip()
                    if optional_columns["item_type"] is not None
                    else ""
                ),
                unit_of_measure=(
                    str(row[optional_columns["unit_of_measure"]] or "case").strip()
                    if optional_columns["unit_of_measure"] is not None
                    else "case"
                ),
                quantity=int(_number(row[columns["quantity"]], row_number, "số lượng")),
                length_m=length,
                width_m=width,
                height_m=height,
                gross_weight_kg=gw_kg,
                dimension_unit=dimension_unit,
            )
        )
    if not items:
        raise ExcelImportError("File Excel không có dòng mặt hàng hợp lệ")

    return {
        "items": items,
        "detected_tenant_code": detected_tenant_code,
        "detected_tenant_name": detected_tenant_name,
        "oversized_items": oversized_items,
    }


def parse_items(workbook_bytes: bytes) -> list[ItemInput]:
    res = parse_excel_inbound(workbook_bytes)
    return res["items"]


def _infer_dimension_unit(dimensions: tuple[float, float, float]) -> str:
    """Infer common spreadsheet units when the workbook has no unit column."""
    if not dimensions:
        return "m"
    max_d = max(dimensions)
    min_d = min(dimensions)

    # Nếu chiều lớn nhất <= 3.0 m, chắc chắn là mét (m)
    if max_d <= 3.0:
        return "m"

    # Nếu chiều lớn nhất > 250 (ví dụ 300, 400, 600, 1200), chắc chắn là milimet (mm)
    if max_d > 250:
        return "mm"

    # Nếu chiều nhỏ nhất < 50 và chiều lớn nhất <= 250 (ví dụ 60x40x30, 120x80x40, 180x60x30 cm), chắc chắn là centimet (cm)
    if min_d < 50:
        return "cm"

    # Nếu cả 3 chiều đều >= 100 (ví dụ 200x150x100), hầu như là mm
    if min_d >= 100:
        return "mm"

    return "cm"
